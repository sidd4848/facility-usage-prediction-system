"""Excel prediction-review workbook: Summary, Predictions vs Actual, Error analysis."""
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config as C

GREEN = PatternFill("solid", fgColor="C6EFCE")
RED = PatternFill("solid", fgColor="FFC7CE")
AMBER = PatternFill("solid", fgColor="FFEB9C")
HEAD = PatternFill("solid", fgColor="1F3864")
G_PAST, G_PRED, G_ACT, G_MATCH = (PatternFill("solid", fgColor=c) for c in ("7F7F7F", "2E75B6", "548235", "7030A0"))
WHITE = Font(color="FFFFFF", bold=True)
BOLD = Font(bold=True)
THIN = Side(style="thin", color="BFBFBF")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
FN = dict(enumerate(C.FACILITY_NAMES))


def _hhmm(minutes):
    m = int(minutes)
    return f"{m // 60:02d}:{m % 60:02d}"


def review_table(m: pd.DataFrame) -> pd.DataFrame:
    d = pd.DataFrame({
        "Record ref": m.target_booking_id,
        "Resident": m.resident_id,
        "Past booking (older)": m.past_1,
        "Past booking (latest)": m.past_2,
        "Pred facility": m.pred_fac.map(FN),
        "Pred day": m.pred_dow.map(dict(enumerate(C.DOW_NAMES))),
        "Pred use time": m.pred_hour.map(lambda h: f"{h:02d}:00"),
        "Pred nudge time": [f"{C.DOW_NAMES[d_]} {_hhmm(t)}" for d_, t in zip(m.pred_nudge_dow, m.pred_nudge_tod_min)],
        "Actual facility": m.act_fac.map(FN),
        "Actual day": m.act_dow.map(dict(enumerate(C.DOW_NAMES))),
        "Actual use time": m.act_hour.map(lambda h: f"{h:02d}:00"),
        "Actual booked at": [f"{C.DOW_NAMES[d_]} {_hhmm(t)}" for d_, t in zip(m.act_bdow, m.act_tod_min)],
        "Facility": m.ok_fac.map({True: "YES", False: "NO"}),
        "Day": m.ok_dow.map({True: "YES", False: "NO"}),
        "Hour": m.ok_hour.map({True: "YES", False: "NO"}),
        "Nudge": m.ok_nudge.map({True: "YES", False: "NO"}),
        "Outputs correct": [f"{s} of 4" for s in m.score],
        "MATCH (all 4)": ["YES" if s == 4 else "NO" for s in m.score],
        "Nudge error (min, time of day)": m.tod_err_min.round(0).astype(int),
        "Actual booking timestamp": m.act_booking_ts.dt.strftime("%Y-%m-%d %H:%M"),
        "Actual usage timestamp": m.act_usage_ts.dt.strftime("%Y-%m-%d %H:%M"),
        "Bookings in history": m.n_hist,
    })
    return d


def _write_df(ws, df, r0, c0=1, title=None, pct_cols=()):
    if title:
        ws.cell(r0, c0, title).font = Font(bold=True, size=12)
        r0 += 1
    for j, col in enumerate(df.columns):
        c = ws.cell(r0, c0 + j, str(col))
        c.font, c.fill, c.border = WHITE, HEAD, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, row in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(row):
            c = ws.cell(r0 + i, c0 + j, v.item() if hasattr(v, "item") else v)
            c.border = BOX
            if isinstance(c.value, float):
                c.number_format = "0.0%" if df.columns[j] in pct_cols else "0.000"
    return r0 + len(df) + 2


def build_workbook(m, metrics, base_metrics, errors, selection, path=C.REVIEW_XLSX):
    wb = Workbook()

    # ---------------- Summary ----------------
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Facility Usage Prediction - holdout test results"
    ws["A1"].font = Font(bold=True, size=14)
    ws["A2"] = (f"Chronological holdout: predictions for resident bookings created after {C.VAL_END[:10]} "
                f"(model fit on bookings <= {C.VAL_END[:10]}; test data never used for fitting or tuning). "
                f"{metrics['n_test_records']:,} test records.")
    ws["A3"] = ("Each record: given the resident's history up to their latest booking, predict the NEXT booking - "
                "facility, usage weekday, usage hour, and nudge time (weekday + time of day to notify = predicted booking time).")
    ws["A4"] = (f"Nudge counts as correct when the weekday matches and the time of day is within +-{C.NUDGE_TOL_MIN} min. "
                "Record MATCH = YES only when all 4 outputs are correct.")
    ov = metrics["overall"]
    rows = [
        ("Facility", "Top-1 accuracy", metrics["facility"]["accuracy"]),
        ("Facility", "Top-3 accuracy", metrics["facility"].get("top3_accuracy")),
        ("Facility", "Macro-F1 (class-imbalance aware)", metrics["facility"]["macro_f1"]),
        ("Usage day", "Exact-weekday accuracy", metrics["usage_day"]["accuracy"]),
        ("Usage day", "Within +-1 day", metrics["usage_day"]["within_1_day"]),
        ("Usage hour", "Exact-hour accuracy", metrics["usage_hour"]["accuracy"]),
        ("Usage hour", "Within +-1 hour", metrics["usage_hour"]["within_1_hour"]),
        ("Usage hour", "MAE (hours)", metrics["usage_hour"]["mae_hours"]),
        ("Nudge time", f"Correct (weekday + time within +-{C.NUDGE_TOL_MIN} min)", metrics["nudge_time"]["accuracy"]),
        ("Nudge time", "Booking-weekday accuracy", metrics["nudge_time"]["booking_weekday_accuracy"]),
        ("Nudge time", "Time-of-day MAE (min)", metrics["nudge_time"]["time_of_day_mae_min"]),
        ("OVERALL", "Average share of the 4 outputs correct", ov["overall_accuracy_pct"] / 100),
        ("OVERALL", "Records with all 4 correct", ov["all_4_correct"]),
        ("OVERALL", "Records with >= 3 of 4 correct", ov["at_least_3_correct"]),
        ("OVERALL", "Records with >= 2 of 4 correct", ov["at_least_2_correct"]),
    ]
    sm = pd.DataFrame(rows, columns=["Output", "Metric", "Model"])
    if base_metrics:
        def pick(bm, o, k):
            if o == "OVERALL":
                return {"Average share of the 4 outputs correct": bm["overall"]["overall_accuracy_pct"] / 100,
                        "Records with all 4 correct": bm["overall"]["all_4_correct"],
                        "Records with >= 3 of 4 correct": bm["overall"]["at_least_3_correct"],
                        "Records with >= 2 of 4 correct": bm["overall"]["at_least_2_correct"]}[k]
            key = {"Top-1 accuracy": ("facility", "accuracy"), "Macro-F1 (class-imbalance aware)": ("facility", "macro_f1"),
                   "Exact-weekday accuracy": ("usage_day", "accuracy"), "Within +-1 day": ("usage_day", "within_1_day"),
                   "Exact-hour accuracy": ("usage_hour", "accuracy"), "Within +-1 hour": ("usage_hour", "within_1_hour"),
                   "MAE (hours)": ("usage_hour", "mae_hours"), "Booking-weekday accuracy": ("nudge_time", "booking_weekday_accuracy"),
                   "Time-of-day MAE (min)": ("nudge_time", "time_of_day_mae_min")}.get(k)
            if k.startswith("Correct (weekday"):
                key = ("nudge_time", "accuracy")
            return bm[key[0]][key[1]] if key else None
        for name, bm in base_metrics.items():
            sm[name] = [pick(bm, o, k) for o, k, _ in rows]
    pct_rows = {i for i, (_, k, _) in enumerate(rows) if "MAE" not in k}
    r0 = 6
    for j, col in enumerate(sm.columns):
        c = ws.cell(r0, 1 + j, col)
        c.font, c.fill, c.border = WHITE, HEAD, BOX
        c.alignment = Alignment(wrap_text=True)
    for i, row in enumerate(sm.itertuples(index=False)):
        for j, v in enumerate(row):
            c = ws.cell(r0 + 1 + i, 1 + j, v)
            c.border = BOX
            if j >= 2 and isinstance(v, float):
                c.number_format = "0.0%" if i in pct_rows else "0.00"
            if row[0] == "OVERALL":
                c.font = BOLD
    r = r0 + len(sm) + 3
    dist = pd.DataFrame({"Outputs correct": [f"{k} of 4" for k in range(5)],
                         "Records": [ov["score_distribution"].get(k, 0) for k in range(5)]})
    dist["Share"] = dist["Records"] / dist["Records"].sum()
    r = _write_df(ws, dist, r, 1, "Score distribution", pct_cols=("Share",))
    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 52
    for k in "CDEF":
        ws.column_dimensions[k].width = 20

    # ---------------- Predictions vs Actual ----------------
    wr = wb.create_sheet("Predictions vs Actual")
    d = review_table(m)
    groups = [("", 1, 2, None), ("PAST BOOKINGS  facility / day / use time / booked at", 3, 4, G_PAST),
              ("PREDICTION  facility / day / use time / nudge time", 5, 8, G_PRED),
              ("ACTUAL  facility / day / use time / booked at", 9, 12, G_ACT),
              ("MATCH", 13, 18, G_MATCH), ("DETAIL", 19, 22, G_PAST)]
    for text, a, b, fill in groups:
        wr.merge_cells(start_row=1, start_column=a, end_row=1, end_column=b)
        c = wr.cell(1, a, text)
        if fill is not None:
            c.fill, c.font = fill, WHITE
        c.alignment = Alignment(horizontal="center")
    for j, col in enumerate(d.columns, start=1):
        c = wr.cell(2, j, col)
        c.font, c.fill, c.border = WHITE, HEAD, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    for i, row in enumerate(d.itertuples(index=False), start=3):
        for j, v in enumerate(row, start=1):
            c = wr.cell(i, j, v.item() if hasattr(v, "item") else v)
            c.border = BOX
            if 13 <= j <= 16 or j == 18:
                c.fill = GREEN if v == "YES" else RED
                c.alignment = Alignment(horizontal="center")
            if j == 17:
                s = int(str(v)[0])
                c.fill = GREEN if s == 4 else (AMBER if s >= 2 else RED)
                c.alignment = Alignment(horizontal="center")
    widths = [11, 10, 34, 34, 13, 9, 11, 14, 13, 10, 11, 14, 9, 7, 7, 8, 11, 12, 13, 18, 18, 10]
    for j, w_ in enumerate(widths, start=1):
        wr.column_dimensions[get_column_letter(j)].width = w_
    wr.row_dimensions[2].height = 32
    wr.freeze_panes = "C3"
    wr.auto_filter.ref = f"A2:{get_column_letter(len(d.columns))}{len(d) + 2}"

    # ---------------- Error analysis ----------------
    we = wb.create_sheet("Error analysis")
    r = 1
    titles = {"by_history_size": "Accuracy by amount of resident history (sparsity)",
              "by_actual_facility": "Accuracy by actual facility (imbalance)",
              "day_hour_given_facility_correctness": "Cascade effect: downstream accuracy when facility was predicted correctly vs wrongly",
              "top_facility_confusions": "Top facility confusions (actual -> predicted)",
              "by_resident_drift": "Residents whose behaviour changed during the history (hidden generator flag, analysis only)",
              "by_lead_archetype": "By resident booking-lead style (hidden generator flag, analysis only)",
              "hour_error_distribution": "Usage-hour error (predicted - actual), clipped at +-3"}
    for k, df in errors.items():
        pc = tuple(c for c in df.columns if c in ("facility", "day", "hour", "nudge", "avg_of_4", "all_4", "share"))
        r = _write_df(we, df, r, 1, titles.get(k, k), pct_cols=tuple(c for c in pc if c != "avg_of_4"))
    we.column_dimensions["A"].width = 34
    for k in "BCDEFGH":
        we.column_dimensions[k].width = 14

    # ---------------- Model selection ----------------
    if selection is not None:
        wm = wb.create_sheet("Validation model comparison")
        wm["A1"] = "Validation-period comparison (stage-wise, teacher-forced). Test data not involved."
        wm["A1"].font = BOLD
        _write_df(wm, selection, 3)
        for k, w_ in zip("ABCDE", (8, 40, 60, 12, 10)):
            wm.column_dimensions[k].width = w_

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path
